#!/usr/bin/env python3
# 체험 기록 집계 JSON(tools/exp-report.mjs 결과)을 엑셀로. 시트: 요약, 단계별 도달, 화면별 머문 시간, 이탈 지점, 버튼, 사람별 기록
import json
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

src, out = sys.argv[1], sys.argv[2]
data = json.load(open(src, encoding='utf-8'))
rep = data.get('report') or {}
groups = list(rep.keys()) or ['전체']
HEAD = Font(bold=True, color='FFFFFF')
FILL = PatternFill('solid', fgColor='1F6FEB')


def sheet(wb, title, header, rows, widths=None):
    ws = wb.create_sheet(title)
    ws.append(header)
    for c in ws[1]:
        c.font = HEAD
        c.fill = FILL
        c.alignment = Alignment(vertical='center')
    for r in rows:
        ws.append(r)
    for i, w in enumerate(widths or [], start=1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w
    ws.freeze_panes = 'A2'
    return ws


wb = Workbook()
wb.remove(wb.active)
pct = lambda x: round((x or 0) * 100, 1)
summary = []
for g in groups:
    r = rep.get(g) or {}
    summary.append([r.get('label', g), r.get('participants', 0), r.get('completed', 0), pct(r.get('completionRate')), r.get('declined', 0), r.get('medianTotalSec'), r.get('withErrors', 0)])
if not rep:
    summary.append(['전체', 0, 0, 0, 0, None, 0])
sheet(wb, '요약', ['출처', '참여(명)', '이송 끝까지(명)', '완수율(%)', '불가 응답(명)', '체험 시간 중앙값(초)', '오류 본 사람(명)'], summary, [16, 10, 14, 10, 12, 18, 14])
ws = wb['요약']
ws.append([])
ws.append([f"만든 시각 {data.get('generatedAt', '')}. 이름, 연락처, 입력한 글은 기록하지 않음. 시험 출처는 빼고 셈"])

rows = []
for g in groups:
    for f in (rep.get(g) or {}).get('funnel', []):
        rows.append([rep[g].get('label', g), f['label'], f['reached'], pct(f['rate'])])
sheet(wb, '단계별 도달', ['출처', '단계', '닿은 사람(명)', '도달률(%)'], rows, [16, 16, 14, 12])

rows = []
for g in groups:
    for d in (rep.get(g) or {}).get('dwell', []):
        rows.append([rep[g].get('label', g), d['label'], d['screen'], d['n'], d['medianSec'], d['maxSec']])
sheet(wb, '화면별 머문 시간', ['출처', '화면', '화면 키', '머문 횟수', '중앙값(초)', '가장 길게(초)'], rows, [16, 24, 14, 10, 12, 14])

rows = []
for g in groups:
    for d in (rep.get(g) or {}).get('dropoff', []):
        rows.append([rep[g].get('label', g), d['label'], d['screen'], d['n']])
sheet(wb, '이탈 지점', ['출처', '마지막으로 본 화면', '화면 키', '사람(명)'], rows, [16, 24, 14, 10])

rows = []
for g in groups:
    for a in (rep.get(g) or {}).get('actions', []):
        rows.append([rep[g].get('label', g), a['key'], a['n']])
sheet(wb, '버튼', ['출처', '버튼 키', '누른 횟수'], rows, [16, 18, 10])

rows = []
for s in data.get('sessions', []):
    rows.append([s['id'], s['src'], s['dev'], s['wb'], s.get('firstAt'), s['totalSec'], '예' if s['done'] else '아니오', s['lastView'], s['errors'], ' > '.join(s['views'])])
sheet(wb, '사람별 기록', ['체험 번호(앞 6자)', '출처', '기기', '화면 폭', '시작(ms)', '체험 시간(초)', '이송 끝', '마지막 화면', '오류', '본 화면 순서'], rows, [16, 10, 9, 8, 15, 12, 9, 12, 6, 80])
wb.save(out)
print('엑셀 저장', out)
